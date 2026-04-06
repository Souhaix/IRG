"""Download / upload an Excel file from SharePoint via Microsoft Graph API.

Write logic uses the excel_mapping.json configuration to find the correct
sheet, row (by label), and column (by date).
"""

import io
import logging
from datetime import date, datetime

import msal
import openpyxl
import requests

from config import SharePointConfig
from excel_mapping import get_spend_to_cell_mapping

logger = logging.getLogger(__name__)

GRAPH_BASE = "https://graph.microsoft.com/v1.0"


# ---------------------------------------------------------------------------
# SharePoint auth & file operations
# ---------------------------------------------------------------------------

def _get_access_token(cfg: SharePointConfig) -> str:
    """Acquire an app-only token via client credentials flow."""
    authority = f"https://login.microsoftonline.com/{cfg.tenant_id}"
    app = msal.ConfidentialClientApplication(
        cfg.client_id,
        authority=authority,
        client_credential=cfg.client_secret,
    )
    result = app.acquire_token_for_client(scopes=["https://graph.microsoft.com/.default"])
    if "access_token" not in result:
        raise RuntimeError(f"SharePoint auth failed: {result.get('error_description')}")
    return result["access_token"]


def _get_site_id(token: str, site_name: str) -> str:
    """Resolve the SharePoint site ID from its name."""
    resp = requests.get(
        f"{GRAPH_BASE}/sites?search={site_name}",
        headers={"Authorization": f"Bearer {token}"},
        timeout=30,
    )
    resp.raise_for_status()
    sites = resp.json().get("value", [])
    if not sites:
        raise RuntimeError(f"SharePoint site '{site_name}' not found")
    return sites[0]["id"]


def _file_url(site_id: str, drive_id: str, file_path: str) -> str:
    if drive_id:
        return f"{GRAPH_BASE}/sites/{site_id}/drives/{drive_id}/root:{file_path}:"
    return f"{GRAPH_BASE}/sites/{site_id}/drive/root:{file_path}:"


def download_workbook(cfg: SharePointConfig) -> openpyxl.Workbook:
    """Download the Excel file from SharePoint and return an openpyxl Workbook."""
    token = _get_access_token(cfg)
    site_id = _get_site_id(token, cfg.site_name)
    url = _file_url(site_id, cfg.drive_id, cfg.file_path) + "/content"

    resp = requests.get(
        url,
        headers={"Authorization": f"Bearer {token}"},
        timeout=60,
    )
    resp.raise_for_status()

    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    logger.info("Downloaded workbook from SharePoint (%d bytes)", len(resp.content))
    return wb


def upload_workbook(cfg: SharePointConfig, wb: openpyxl.Workbook) -> None:
    """Upload the modified workbook back to SharePoint."""
    token = _get_access_token(cfg)
    site_id = _get_site_id(token, cfg.site_name)
    url = _file_url(site_id, cfg.drive_id, cfg.file_path) + "/content"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    resp = requests.put(
        url,
        headers={
            "Authorization": f"Bearer {token}",
            "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        },
        data=buf.read(),
        timeout=60,
    )
    resp.raise_for_status()
    logger.info("Uploaded workbook to SharePoint")


# ---------------------------------------------------------------------------
# Excel writing logic
# ---------------------------------------------------------------------------

def _find_date_column(ws, target_date: date, mapping: dict) -> int | None:
    """Find column that matches *target_date* in the date row.

    The date row/location is defined in mapping["dates"]["location"].
    Currently supports "row_1" (dates across row 1).
    """
    date_location = mapping.get("dates", {}).get("location", "row_1")

    if date_location == "row_1":
        return _find_date_in_row(ws, row=1, target_date=target_date)

    # Fallback: assume row_1
    logger.warning("Unknown date location '%s', falling back to row 1", date_location)
    return _find_date_in_row(ws, row=1, target_date=target_date)


def _find_date_in_row(ws, row: int, target_date: date) -> int | None:
    """Scan a row for a cell matching *target_date*."""
    for col in range(1, ws.max_column + 1):
        cell_value = ws.cell(row=row, column=col).value
        if cell_value is None:
            continue
        if _cell_matches_date(cell_value, target_date):
            return col
    return None


def _cell_matches_date(cell_value, target_date: date) -> bool:
    """Check if a cell value matches the target date."""
    # datetime object
    if isinstance(cell_value, datetime):
        return cell_value.date() == target_date
    # date object
    if isinstance(cell_value, date) and not isinstance(cell_value, datetime):
        return cell_value == target_date
    # String — try common formats
    s = str(cell_value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y"):
        try:
            if datetime.strptime(s, fmt).date() == target_date:
                return True
        except ValueError:
            continue
    return False


def _find_label_row(ws, label: str, label_col: int) -> int | None:
    """Find the row where the label column contains *label*."""
    for row in range(1, ws.max_row + 1):
        cell_value = ws.cell(row=row, column=label_col).value
        if cell_value is not None and str(cell_value).strip() == label:
            return row
    return None


def _col_letter(col: int) -> str:
    """Convert 1-based column number to Excel letter (1→A, 27→AA)."""
    result = ""
    while col > 0:
        col, remainder = divmod(col - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _label_col_index(mapping: dict) -> int:
    """Get the 1-based column index for the labels column.

    Reads mapping["labels"]["column"] which should be a letter like "A".
    """
    col_letter = mapping.get("labels", {}).get("column", "A").upper().strip()
    # Convert letter(s) to 1-based index
    index = 0
    for ch in col_letter:
        index = index * 26 + (ord(ch) - ord("A") + 1)
    return index


def write_spend_to_workbook(
    wb: openpyxl.Workbook,
    spend_data: dict[str, float],
    target_date: date,
    mapping: dict,
) -> None:
    """Write spend values into the correct sheet/cell of the workbook.

    Uses the mapping config to:
      1. Determine which sheet and label correspond to each spend key
      2. Find the row by matching the label in the labels column
      3. Find the column by matching the date
    """
    cell_mapping = get_spend_to_cell_mapping(mapping)
    label_col = _label_col_index(mapping)

    for key, spend in spend_data.items():
        target = cell_mapping.get(key)
        if target is None:
            logger.warning("No Excel mapping for spend key '%s' — skipping", key)
            continue

        sheet_name = target["sheet"]
        label = target["label"]

        if not sheet_name or not label:
            logger.warning("Incomplete mapping for '%s' (sheet='%s', label='%s') — skipping", key, sheet_name, label)
            continue

        if sheet_name not in wb.sheetnames:
            logger.warning("Sheet '%s' not found in workbook — skipping '%s'", sheet_name, key)
            continue

        ws = wb[sheet_name]

        # Find the row by label
        row = _find_label_row(ws, label, label_col)
        if row is None:
            logger.warning(
                "Label '%s' not found in column %s of sheet '%s' — skipping '%s'",
                label, _col_letter(label_col), sheet_name, key,
            )
            continue

        # Find the column by date
        col = _find_date_column(ws, target_date, mapping)
        if col is None:
            logger.warning(
                "Date %s not found in sheet '%s' — skipping '%s'",
                target_date, sheet_name, key,
            )
            continue

        ws.cell(row=row, column=col, value=round(spend, 2))
        logger.info(
            "Wrote %.2f → %s!%s%d (key=%s)",
            spend, sheet_name, _col_letter(col), row, key,
        )
